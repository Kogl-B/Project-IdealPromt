import openpyxl

wb = openpyxl.load_workbook('/workspaces/Project-IdealPromt/письма срочно!!!.xlsx')
print('Листы:', wb.sheetnames)
ws = wb.active
print(f'Активный лист: {ws.title}')
print(f'Строк: {ws.max_row}, Колонок: {ws.max_column}')
print('\nДанные из первых 10 строк:')
for row in range(1, min(11, ws.max_row + 1)):
    row_data = []
    for col in range(1, min(6, ws.max_column + 1)):
        cell_value = ws.cell(row=row, column=col).value
        row_data.append(str(cell_value) if cell_value else '')
    print(f'Строка {row}: {row_data}')
