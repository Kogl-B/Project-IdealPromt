#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Скрипт для генерации писем на основе шаблонов Word и данных из Excel."""

import os
import re
from docx import Document
import openpyxl

BASE_DIR = '/workspaces/Project-IdealPromt'
EXCEL_FILE = os.path.join(BASE_DIR, 'письма срочно!!!.xlsx')
TEMPLATE_URGENT = os.path.join(BASE_DIR, 'Письма  УВЕДОМЛЕНИЕ, ПРИГЛАШЕНИЕ.docx')
TEMPLATE_NOT_OCCUPIED = os.path.join(BASE_DIR, 'Письма НЕ ЗАНЯТЫЕ.docx')
OUTPUT_DIR = os.path.join(BASE_DIR, 'Готовые письма')

TEMPLATE_NAME = 'Чуль С.В.'
TEMPLATE_DATE = '30.06.1981'
TEMPLATE_ADDRESS = 'ул. Берестянская, 11-269'

def full_name_to_initials(full_name):
    if not full_name:
        return ''
    full_name = full_name.replace(' null', '').replace('null ', '').replace('null', '')
    parts = full_name.strip().split()
    if len(parts) == 0:
        return ''
    surname = parts[0].capitalize()
    if len(parts) == 1:
        return surname
    elif len(parts) == 2:
        return f"{surname} {parts[1][0].upper()}."
    else:
        return f"{surname} {parts[1][0].upper()}.{parts[2][0].upper()}."

def format_date(date_value):
    if not date_value:
        return ''
    if hasattr(date_value, 'strftime'):
        return date_value.strftime('%d.%m.%Y')
    date_str = str(date_value)
    if ' ' in date_str:
        date_str = date_str.split(' ')[0]
    if '-' in date_str:
        parts = date_str.split('-')
        if len(parts) == 3:
            return f"{parts[2]}.{parts[1]}.{parts[0]}"
    return date_str

def read_excel_data(filepath):
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    data = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=1).value
        date = ws.cell(row=row, column=2).value
        address = ws.cell(row=row, column=3).value
        if not name:
            continue
        full_name = str(name).strip()
        full_name_clean = full_name.replace(' null', '').replace('null ', '').replace('null', '').strip()
        data.append({
            'name_full': full_name_clean,
            'name_initials': full_name_to_initials(full_name),
            'date': format_date(date),
            'address': str(address).strip() if address else ''
        })
    return data

def replace_text_in_runs(paragraph, old_text, new_text):
    replaced = False
    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            replaced = True
    if replaced:
        return True
    if old_text in paragraph.text:
        full_text = paragraph.text
        new_full_text = full_text.replace(old_text, new_text)
        if paragraph.runs:
            paragraph.runs[0].text = new_full_text
            for run in paragraph.runs[1:]:
                run.text = ''
            return True
    return False

def replace_in_document(doc, replacements):
    replaced_count = {}
    for key in replacements.keys():
        replaced_count[key] = 0
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text and old_text in paragraph.text:
                if replace_text_in_runs(paragraph, old_text, new_text):
                    replaced_count[old_text] += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old_text, new_text in replacements.items():
                        if old_text and old_text in paragraph.text:
                            if replace_text_in_runs(paragraph, old_text, new_text):
                                replaced_count[old_text] += 1
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for paragraph in header.paragraphs:
                    for old_text, new_text in replacements.items():
                        if old_text and old_text in paragraph.text:
                            if replace_text_in_runs(paragraph, old_text, new_text):
                                replaced_count[old_text] += 1
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for paragraph in footer.paragraphs:
                    for old_text, new_text in replacements.items():
                        if old_text and old_text in paragraph.text:
                            if replace_text_in_runs(paragraph, old_text, new_text):
                                replaced_count[old_text] += 1
    return replaced_count

def generate_urgent_letter(template_path, output_path, data):
    doc = Document(template_path)
    replacements = {
        TEMPLATE_NAME: data['name_initials'],
        TEMPLATE_DATE: data['date'],
        TEMPLATE_ADDRESS: data['address'],
    }
    replaced = replace_in_document(doc, replacements)
    doc.save(output_path)
    print(f"  Создано: {os.path.basename(output_path)}")
    print(f"    Замены: Имя={replaced[TEMPLATE_NAME]}, Дата={replaced[TEMPLATE_DATE]}, Адрес={replaced[TEMPLATE_ADDRESS]}")

def generate_not_occupied_letter(template_path, output_path, data):
    doc = Document(template_path)
    replacements = {
        TEMPLATE_NAME: data['name_initials'],
        TEMPLATE_ADDRESS: data['address'],
    }
    replaced = replace_in_document(doc, replacements)
    doc.save(output_path)
    print(f"  Создано: {os.path.basename(output_path)}")
    print(f"    Замены: Имя={replaced[TEMPLATE_NAME]}, Адрес={replaced[TEMPLATE_ADDRESS]}")

def main():
    print("=" * 60)
    print("Генератор писем на основе шаблонов")
    print("=" * 60)
    for filepath, name in [(EXCEL_FILE, 'Excel с данными'), 
                           (TEMPLATE_URGENT, 'Шаблон уведомления'),
                           (TEMPLATE_NOT_OCCUPIED, 'Шаблон "не занятые"')]:
        if not os.path.exists(filepath):
            print(f"ОШИБКА: Файл не найден: {filepath}")
            return
        print(f"✓ Найден: {name}")
    
    data_list = read_excel_data(EXCEL_FILE)
    print(f"\nНайдено записей для обработки: {len(data_list)}")
    if not data_list:
        print("ОШИБКА: Нет данных для обработки!")
        return
    
    print("\n=== Пример преобразования данных (первые 3 записи) ===")
    for i, d in enumerate(data_list[:3], 1):
        print(f"  {i}. {d['name_full']}")
        print(f"     -> {d['name_initials']}, дата: {d['date']}, адрес: {d['address'][:40]}...")
    
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f"\nПапка для результатов: {OUTPUT_DIR}")
    
    TEST_MODE = True
    if TEST_MODE:
        print("\n*** ТЕСТОВЫЙ РЕЖИМ - обрабатывается только 1 запись ***")
        data_list = data_list[:1]
    
    print("\n--- Генерация писем ---")
    for i, data in enumerate(data_list, 1):
        print(f"\nОбработка записи {i}:")
        print(f"  Полное имя: {data['name_full']}")
        print(f"  Инициалы: {data['name_initials']}")
        print(f"  Дата: {data['date']}")
        print(f"  Адрес: {data['address']}")
        
        safe_name = re.sub(r'[<>:"/\\|?*]', '_', data['name_initials'])
        
        output_urgent = os.path.join(OUTPUT_DIR, f"Уведомление_{safe_name}.docx")
        generate_urgent_letter(TEMPLATE_URGENT, output_urgent, data)
        
        output_not_occupied = os.path.join(OUTPUT_DIR, f"Не_занятые_{safe_name}.docx")
        generate_not_occupied_letter(TEMPLATE_NOT_OCCUPIED, output_not_occupied, data)
    
    print("\n" + "=" * 60)
    print(f"Готово! Создано писем: {len(data_list) * 2}")
    print(f"Результаты в папке: {OUTPUT_DIR}")
    print("=" * 60)
    
    print("\n=== Проверка созданных документов ===")
    for filename in sorted(os.listdir(OUTPUT_DIR)):
        if filename.endswith('.docx') and not filename.startswith('~'):
            filepath = os.path.join(OUTPUT_DIR, filename)
            print(f"\n--- Файл: {filename} ---")
            doc = Document(filepath)
            found_key_info = False
            for i, para in enumerate(doc.paragraphs[:30], 1):
                text = para.text.strip()
                if text:
                    if any(x in text.lower() for x in ['чуль', 'берестянская', '30.06.1981', 'гражданин', 'проживающ']):
                        print(f"  >>> [{i}] {text[:150]}{'...' if len(text) > 150 else ''}")
                        found_key_info = True
            if not found_key_info:
                print("  ⚠️ Ключевые строки с данными не найдены!")

if __name__ == '__main__':
    main()
