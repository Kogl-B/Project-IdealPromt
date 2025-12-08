#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from docx import Document

# Анализируем созданный документ
doc = Document('/workspaces/Project-IdealPromt/Готовые письма/Уведомление_Аникович В.Н..docx')

print("=== ПОЛНЫЙ АНАЛИЗ ДОКУМЕНТА УВЕДОМЛЕНИЕ ===\n")

for i, para in enumerate(doc.paragraphs, 1):
    text = para.text
    # Показываем все параграфы, которые могут содержать данные для замены
    if any(keyword in text for keyword in ['гражданин', 'проживающ', 'Берестянская', 'адрес', 'рождения', '1969', '1981', 'Чуль', 'Аникович', 'АНИКОВИЧ']):
        print(f"[{i}] {text}")
        print()

print("\n=== АНАЛИЗ ДОКУМЕНТА НЕ ЗАНЯТЫЕ ===\n")
doc2 = Document('/workspaces/Project-IdealPromt/Готовые письма/Не_занятые_Аникович В.Н..docx')

for i, para in enumerate(doc2.paragraphs, 1):
    text = para.text
    if text.strip():
        print(f"[{i}] {text}")
