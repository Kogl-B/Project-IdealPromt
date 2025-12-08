#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Анализ структуры Word документа"""

from docx import Document

doc = Document('/workspaces/Project-IdealPromt/Письма  УВЕДОМЛЕНИЕ, ПРИГЛАШЕНИЕ.docx')

print("=== Все параграфы документа ===")
for i, para in enumerate(doc.paragraphs):
    text = para.text
    if text.strip():
        print(f"\n[{i}]: {repr(text)}")

print("\n\n=== Секции документа ===")
print(f"Количество секций: {len(doc.sections)}")

print("\n\n=== Таблицы ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\nТаблица {t_idx}:")
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            if cell.text.strip():
                print(f"  Ячейка [{r_idx},{c_idx}]: {repr(cell.text[:200])}")

# Поиск конкретных строк
print("\n\n=== Поиск ключевых фраз ===")
search_terms = ['Берестянская', 'Чуль', 'АНИКОВИЧ', 'гражданину', 'проживающему', '30.06.1981']
for para in doc.paragraphs:
    for term in search_terms:
        if term.lower() in para.text.lower():
            print(f"\nНайдено '{term}':")
            print(f"  {repr(para.text[:300])}")
            break
